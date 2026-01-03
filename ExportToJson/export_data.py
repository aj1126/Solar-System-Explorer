import openpyxl
import json

# Configuration
input_file = 'Solar System Explorer v1.4.xlsm'
output_file = 'solar_system_data.json'

# Color Mapping (Based on your Color Key screenshot)
# We can map specific planets to their colors. Others will get a default.
color_map = {
    "Mercury": "#A9A9A9",
    "Venus": "#FFC966",
    "Earth": "#4DA6FF",
    "Mars": "#FF9933",
    "Jupiter": "#FF9966",
    "Saturn": "#FFD966",
    "Uranus": "#66CCCC",
    "Neptune": "#3366FF",
    "Pluto": "#9966CC",
    "Ceres": "#9966CC",
    "Eris": "#9966CC",
    "Haumea": "#9966CC",
    "Makemake": "#9966CC",
    "Gonggong": "#9966CC",
    "Quaoar": "#66CCFF",
    "Sedna": "#999999",
    "Orcus": "#66CCFF",
    "Salacia": "#66CCFF",
    "Vesta": "#CCCC99" 
}
default_color = "#CCCCCC" # Grey for asteroids/others

def get_color(name):
    return color_map.get(name, default_color)

print(f"Loading {input_file}...")
try:
    # Load workbook, data_only=True ensures we get values, not formulas
    wb = openpyxl.load_workbook(input_file, data_only=True, keep_vba=False)
    
    # We target 'Sorting Data' because your screenshot shows it has clean 
    # numbers for SMA_AU and e_value, whereas the main sheet might have text like "1.00 AU".
    if "Sorting Data" not in wb.sheetnames:
        raise ValueError("Sheet 'Sorting Data' not found!")
    
    ws = wb["Sorting Data"]
    
    # Locate columns based on your screenshot of 'Sorting Data'
    # Headers are likely in Row 1: 
    # Name (A), Satellite_# (B), Radius_km (C), SMA_AU (D), e_value (E)
    
    data_list = []
    
    # Iterate through rows, skipping the header (start at row 2)
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[0]  # Column A
        
        # Stop if we hit an empty row
        if not name:
            break
            
        radius = row[2] # Column C
        sma = row[3]    # Column D (Semi-Major Axis)
        ecc = row[4]    # Column E (Eccentricity)
        
        # Basic validation to ensure we have numbers
        if isinstance(sma, (int, float)) and isinstance(ecc, (int, float)):
            planet_obj = {
                "name": str(name).strip(),
                "radius_km": radius,
                "a": float(sma),      # Semi-Major Axis in AU
                "e": float(ecc),      # Eccentricity
                "color": get_color(name)
            }
            data_list.append(planet_obj)

    # Save to JSON
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(data_list, f, indent=2)
        
    print(f"Success! Exported {len(data_list)} objects to {output_file}")

except Exception as e:
    print(f"Error: {e}")